from __future__ import annotations

import re
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[1]
DESIGN_DIR = ROOT / "docs" / "design"
ASSET_DIR = ROOT / ".docx-assets"
ASSET_DIR.mkdir(exist_ok=True)

NAVY = "17365D"
BLUE = "2F75B5"
PALE = "EAF2F8"
GRID = "D9E2F3"
TEXT = "1F2937"

SOURCES = [
    DESIGN_DIR / "Lead_Engine_Technical_Design_v2.md",
    DESIGN_DIR / "Delivery_Factory_Technical_Design_v1.md",
    DESIGN_DIR / "Platform_Integration_and_Operations_Design_v1.md",
]


def font(size: int, bold: bool = False):
    try:
        return ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", size)
    except OSError:
        return ImageFont.load_default()


def draw_box(draw, xy, label, fill="#EAF2F8", outline="#2F75B5"):
    x1, y1, x2, y2 = xy
    draw.rounded_rectangle(xy, radius=14, fill=fill, outline=outline, width=3)
    words = label.split()
    lines, current = [], ""
    for word in words:
        candidate = f"{current} {word}".strip()
        if draw.textlength(candidate, font=font(23)) <= (x2 - x1 - 30):
            current = candidate
        else:
            lines.append(current)
            current = word
    if current:
        lines.append(current)
    line_h = 30
    start_y = (y1 + y2 - line_h * len(lines)) / 2
    for idx, line in enumerate(lines):
        width = draw.textlength(line, font=font(23))
        draw.text(((x1 + x2 - width) / 2, start_y + idx * line_h), line, fill="#17365D", font=font(23))


def arrow(draw, start, end):
    draw.line([start, end], fill="#5B6B7C", width=4)
    ex, ey = end
    sx, sy = start
    if abs(ex - sx) >= abs(ey - sy):
        direction = 1 if ex > sx else -1
        pts = [(ex, ey), (ex - 14 * direction, ey - 9), (ex - 14 * direction, ey + 9)]
    else:
        direction = 1 if ey > sy else -1
        pts = [(ex, ey), (ex - 9, ey - 14 * direction), (ex + 9, ey - 14 * direction)]
    draw.polygon(pts, fill="#5B6B7C")


def architecture_diagram(kind: str, index: int) -> Path:
    output = ASSET_DIR / f"{kind}-{index}.png"
    image = Image.new("RGB", (1600, 660), "white")
    draw = ImageDraw.Draw(image)
    if kind == "lead":
        boxes = [
            ((70, 70, 470, 190), "Next.js Lead workspace"),
            ((600, 70, 1000, 190), "Lead domain services"),
            ((1130, 70, 1530, 190), "Lead Supabase project"),
            ((70, 390, 470, 510), "Crawlee adapters"),
            ((600, 390, 1000, 510), "Trigger.dev workflows"),
            ((1130, 390, 1530, 510), "Delivery API v1"),
        ]
        for xy, label in boxes:
            draw_box(draw, xy, label)
        arrow(draw, (470, 130), (600, 130)); arrow(draw, (1000, 130), (1130, 130))
        arrow(draw, (470, 450), (600, 450)); arrow(draw, (800, 390), (800, 190))
        arrow(draw, (1000, 450), (1130, 450))
    elif kind == "delivery" and index == 1:
        boxes = [
            ((70, 70, 430, 190), "Signed handoff API"),
            ((620, 70, 980, 190), "Verification gate"),
            ((1170, 70, 1530, 190), "Delivery domain"),
            ((70, 390, 430, 510), "Delivery Supabase"),
            ((620, 390, 980, 510), "Trigger.dev workflows"),
            ((1170, 390, 1530, 510), "Tool adapters"),
        ]
        for xy, label in boxes:
            draw_box(draw, xy, label, fill="#EAF6EF", outline="#287A4B")
        arrow(draw, (430, 130), (620, 130)); arrow(draw, (980, 130), (1170, 130))
        arrow(draw, (1350, 190), (1350, 390)); arrow(draw, (1170, 450), (980, 450))
        arrow(draw, (620, 450), (430, 450)); arrow(draw, (250, 390), (1170, 170))
    elif kind == "delivery":
        labels = ["Intake", "Onboarding", "Baseline", "Planned", "Building", "Verifying", "Release", "Handover", "Closed"]
        x, y = 70, 100
        positions = []
        for i, label in enumerate(labels):
            row, col = divmod(i, 5)
            bx = x + col * 300
            by = y + row * 260
            positions.append((bx, by))
            draw_box(draw, (bx, by, bx + 220, by + 95), label, fill="#EAF6EF", outline="#287A4B")
        for i in range(4):
            arrow(draw, (positions[i][0] + 220, positions[i][1] + 47), (positions[i+1][0], positions[i+1][1] + 47))
        arrow(draw, (positions[4][0] + 110, positions[4][1] + 95), (positions[5][0] + 110, positions[5][1]))
        for i in range(5, 8):
            arrow(draw, (positions[i][0] + 220, positions[i][1] + 47), (positions[i+1][0], positions[i+1][1] + 47))
    elif kind == "platform" and index == 1:
        boxes = [
            ((80, 70, 500, 190), "Lead Engine lead.raphah.io"),
            ((1100, 70, 1520, 190), "Delivery Factory delivery.raphah.io"),
            ((80, 410, 500, 530), "Lead Supabase project"),
            ((1100, 410, 1520, 530), "Delivery Supabase project"),
        ]
        for xy, label in boxes:
            draw_box(draw, xy, label)
        arrow(draw, (500, 115), (1100, 115)); arrow(draw, (1100, 150), (500, 150))
        arrow(draw, (290, 190), (290, 410)); arrow(draw, (1310, 190), (1310, 410))
        draw.text((680, 75), "Signed handoff v1", fill="#17365D", font=font(22))
        draw.text((675, 165), "Signed feedback v1", fill="#17365D", font=font(22))
    else:
        labels = ["Lead transaction", "Lead outbox", "Delivery API", "Delivery inbox", "Receipt and project"]
        for i, label in enumerate(labels):
            bx = 55 + i * 305
            draw_box(draw, (bx, 235, bx + 235, 355), label)
            if i < len(labels) - 1:
                arrow(draw, (bx + 235, 295), (bx + 305, 295))
    image.save(output, dpi=(180, 180))
    return output


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=110, bottom=100, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def set_cant_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    run._r.addnext(fld)


def clean_inline(text: str) -> str:
    text = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1 (\2)", text)
    text = text.replace("**", "").replace("__", "")
    return text.replace("`", "")


def add_table(doc: Document, rows: list[list[str]]):
    if not rows:
        return
    cols = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for r_idx, row in enumerate(rows):
        for c_idx in range(cols):
            cell = table.cell(r_idx, c_idx)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)
            text = clean_inline(row[c_idx]) if c_idx < len(row) else ""
            cell.text = text
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(2)
                paragraph.paragraph_format.space_before = Pt(2)
                for run in paragraph.runs:
                    run.font.name = "Arial"
                    run.font.size = Pt(8.7)
                    run.font.color.rgb = RGBColor.from_string("FFFFFF" if r_idx == 0 else TEXT)
                    run.bold = r_idx == 0
            if r_idx == 0:
                set_cell_shading(cell, NAVY)
            elif r_idx % 2 == 0:
                set_cell_shading(cell, "F4F8FC")
    set_repeat_table_header(table.rows[0])
    for row in table.rows:
        set_cant_split(row)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def configure_styles(doc: Document):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(TEXT)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.12
    for name, size, before, after in [
        ("Title", 28, 0, 14),
        ("Heading 1", 18, 16, 7),
        ("Heading 2", 14, 14, 6),
        ("Heading 3", 11.5, 10, 4),
    ]:
        style = styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
    title_ppr = styles["Title"]._element.get_or_add_pPr()
    border = title_ppr.find(qn("w:pBdr"))
    if border is not None:
        title_ppr.remove(border)


def add_cover(doc: Document, title: str, subtitle: str):
    doc.add_paragraph().paragraph_format.space_after = Pt(80)
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title_lines = {
        "Delivery Factory Technical Design Version 1": [
            "Delivery Factory",
            "Technical Design Version 1",
        ],
        "Platform Integration and Operations Design Version 1": [
            "Platform Integration and Operations",
            "Design Version 1",
        ],
    }.get(title, [title])
    for index, line in enumerate(title_lines):
        title_run = p.add_run(line)
        title_run.font.size = Pt(24 if len(title_lines) > 1 else 28)
        if index < len(title_lines) - 1:
            title_run.add_break()
    p2 = doc.add_paragraph()
    p2.paragraph_format.space_after = Pt(36)
    r2 = p2.add_run(subtitle)
    r2.font.name = "Arial"
    r2.font.size = Pt(16)
    r2.font.color.rgb = RGBColor.from_string(NAVY)
    meta = doc.add_table(rows=4, cols=2)
    meta.alignment = WD_TABLE_ALIGNMENT.LEFT
    values = [
        ("Prepared for", "Raphah"),
        ("Status", "Technical baseline"),
        ("Date", date.today().strftime("%d %B %Y")),
        ("Decision", "Separate product with contract-only integration"),
    ]
    for i, (label, value) in enumerate(values):
        meta.cell(i, 0).text = label
        meta.cell(i, 1).text = value
        set_cell_shading(meta.cell(i, 0), PALE)
        for cell in meta.rows[i].cells:
            set_cell_margins(cell, 110, 130, 110, 130)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            for run in cell.paragraphs[0].runs:
                run.font.name = "Arial"
                run.font.size = Pt(10)
                run.font.color.rgb = RGBColor.from_string(TEXT)
        meta.cell(i, 0).paragraphs[0].runs[0].bold = True
    doc.add_page_break()


def add_contents(doc: Document, headings: list[str]):
    doc.add_heading("Contents", level=1)
    for heading in headings:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.15)
        p.add_run(heading)
    doc.add_page_break()


def build(source: Path):
    text = source.read_text(encoding="utf-8")
    lines = text.splitlines()
    title = lines[0].lstrip("# ").strip()
    subtitle_map = {
        "Lead_Engine_Technical_Design_v2.md": "Independent opportunity management and controlled delivery handoff",
        "Delivery_Factory_Technical_Design_v1.md": "Independent client delivery automation and governance",
        "Platform_Integration_and_Operations_Design_v1.md": "Versioned product boundary and independent operations",
    }
    kind = "lead" if source.name.startswith("Lead") else "delivery" if source.name.startswith("Delivery") else "platform"
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)
    configure_styles(doc)
    add_cover(doc, title, subtitle_map[source.name])
    headings = [clean_inline(line[3:].strip()) for line in lines if line.startswith("## ") and not line.startswith("### ")]
    add_contents(doc, headings)

    in_code = False
    code_type = ""
    code_lines: list[str] = []
    table_rows: list[list[str]] = []
    diagram_index = 0

    def flush_table():
        nonlocal table_rows
        if table_rows:
            usable = [row for row in table_rows if not all(re.fullmatch(r"-+", item.replace(":", "")) for item in row)]
            add_table(doc, usable)
            table_rows = []

    def flush_code():
        nonlocal code_lines, diagram_index
        if not code_lines and code_type != "mermaid":
            return
        if code_type == "mermaid":
            diagram_index += 1
            image_path = architecture_diagram(kind, diagram_index)
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(str(image_path), width=Inches(6.65))
            cap = doc.add_paragraph(f"Figure {diagram_index}  {title} architecture view")
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cap.paragraph_format.space_after = Pt(9)
            for run in cap.runs:
                run.font.size = Pt(9)
                run.italic = True
        else:
            for line in code_lines:
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.25)
                p.paragraph_format.space_after = Pt(1)
                r = p.add_run(line)
                r.font.name = "Liberation Mono"
                r.font.size = Pt(8.4)
        code_lines = []

    for raw in lines[1:]:
        line = raw.rstrip()
        if line.startswith("```"):
            flush_table()
            if not in_code:
                in_code = True
                code_type = line[3:].strip()
                code_lines = []
            else:
                flush_code()
                in_code = False
                code_type = ""
            continue
        if in_code:
            code_lines.append(line)
            continue
        if line.startswith("|") and line.endswith("|"):
            table_rows.append([cell.strip() for cell in line.strip("|").split("|")])
            continue
        flush_table()
        if not line:
            continue
        if line.startswith("### "):
            doc.add_heading(clean_inline(line[4:]), level=3)
        elif line.startswith("## "):
            doc.add_heading(clean_inline(line[3:]), level=1)
        elif line.startswith("# "):
            continue
        elif re.match(r"^\d+\.\s", line):
            match = re.match(r"^(\d+)\.\s+(.*)$", line)
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.25)
            p.paragraph_format.first_line_indent = Inches(-0.25)
            p.add_run(f"{match.group(1)}.  {clean_inline(match.group(2))}")
        elif line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            p.add_run(clean_inline(line[2:]))
        else:
            p = doc.add_paragraph()
            p.add_run(clean_inline(line))
    flush_table()
    if in_code:
        flush_code()

    for section in doc.sections:
        header = section.header.paragraphs[0]
        header.text = title
        header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        for run in header.runs:
            run.font.name = "Arial"
            run.font.size = Pt(8)
            run.font.color.rgb = RGBColor.from_string("667085")
        footer = section.footer.paragraphs[0]
        add_page_number(footer)
        for run in footer.runs:
            run.font.name = "Arial"
            run.font.size = Pt(8)
            run.font.color.rgb = RGBColor.from_string("667085")

    output = source.with_suffix(".docx")
    doc.save(output)
    print(output)


if __name__ == "__main__":
    for source in SOURCES:
        build(source)
